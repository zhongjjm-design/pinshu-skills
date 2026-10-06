"""Actual python-pptx -> LibreOffice -> PNG tests, including hidden slides and geometry."""
from pathlib import Path
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "pinshu-visual-system/scripts"))
from render_editable import render, native_inventory, verify_render
from publish_image import reviewed_editable_source
try:
    from pptx import Presentation
    from pptx.util import Inches
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.dml.color import RGBColor
    from pptx.chart.data import CategoryChartData
except ImportError:
    Presentation = None

AVAILABLE = Presentation is not None and all(shutil.which(n) for n in ("soffice", "pdftoppm", "pdfinfo", "magick"))

if (os.environ.get("CI") == "true" or os.environ.get("PINSHU_REQUIRE_NATIVE_RENDER") == "1") and not AVAILABLE:
    raise RuntimeError("Real native PPTX integration dependencies are required; this run cannot skip them")

@unittest.skipUnless(AVAILABLE, "Install python-pptx and documented native-render dependencies")
class NativePptxRendering(unittest.TestCase):
    def font(self):
        for p in ("/System/Library/Fonts/STHeiti Light.ttc", "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"):
            if Path(p).is_file():
                return Path(p)
        if os.environ.get("CI") == "true" or os.environ.get("PINSHU_REQUIRE_NATIVE_RENDER") == "1":
            raise RuntimeError("The required CJK font is missing")
        self.skipTest("Install the documented CJK font")

    def color_bounds(self, image, color):
        return subprocess.check_output(["magick", str(image), "-fuzz", "1%", "-fill", "black", "+opaque", color,
            "-fill", "white", "-opaque", color, "-trim", "-format", "%wx%h", "info:"], text=True)

    def deck(self, path, width=10, height=7.5, hidden=()):
        prs = Presentation(); prs.slide_width = Inches(width); prs.slide_height = Inches(height)
        for i, color in enumerate(((255, 0, 0), (0, 0, 255), (0, 255, 0))):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2), Inches(2), Inches(3), Inches(3))
            shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor(*color); shape.line.fill.background()
            slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(4), Inches(1)).text_frame.text = f"Original slide {i+1}"
            if i+1 in hidden:
                slide._element.set("show", "0")
        prs.save(path)

    def test_hidden_first_selects_original_second(self):
        self.check_hidden((1,), 2, "blue")

    def test_hidden_middle_selects_original_third(self):
        self.check_hidden((2,), 3, "lime")

    def test_selected_hidden_slide_is_rendered(self):
        self.check_hidden((2,), 2, "blue")

    def check_hidden(self, hidden, selected, color):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); native = root / "deck.pptx"; image = root / "image.png"
            self.deck(native, hidden=hidden)
            receipt = render(native, self.font(), "wechat-article", image, slide=selected)
            self.assertEqual(receipt["slide_mapping"]["pdf_page_count"], 3)
            self.assertIn(f"Original slide {selected}", receipt["native_text"])
            w, h = map(int, self.color_bounds(image, color).split("x"))
            self.assertGreater(w, 300); self.assertLessEqual(abs(w-h), 1)
            self.assertEqual(verify_render(native, image, receipt, "wechat-article")["re_render_pixel_difference_ae"], 0)
            old = dict(receipt, schema="native-render-v1")
            with self.assertRaisesRegex(ValueError, "receipt"):
                verify_render(native, image, old, "wechat-article")

    def test_square_preserves_aspect_for_standard_portrait_and_wide(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for width, height in ((10, 7.5), (7.5, 10), (13.333333, 7.5)):
                with self.subTest(ratio=(width, height)):
                    native, image = root / "deck.pptx", root / "image.png"
                    self.deck(native, width, height)
                    render(native, self.font(), "wechat-article", image, slide=2)
                    w, h = map(int, self.color_bounds(image, "blue").split("x"))
                    self.assertGreater(w, 200); self.assertLessEqual(abs(w-h), 1)

    def test_quoted_percentage_suffix_does_not_multiply_cache(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "literal.pptx"
            prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
            data = CategoryChartData(); data.categories = ["Revenue"]; data.add_series("Growth", [30], number_format='0"%"')
            chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(8), Inches(5), data).chart
            chart.plots[0].has_data_labels = True; chart.plots[0].data_labels.number_format = '0"%"'
            prs.save(path)
            native = native_inventory(path)
            self.assertEqual(str(native["chart_points"][0][1]), "30")
            self.assertEqual(native["formatted_chart_points"], [])

    def test_native_percentage_cache_matches_displayed_percentage(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); native, image = root / "chart.pptx", root / "image.png"
            prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
            data = CategoryChartData(); data.categories = ["Revenue"]; data.add_series("Growth", [0.3], number_format="0%")
            chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(8), Inches(5), data).chart
            chart.plots[0].has_data_labels = True; chart.plots[0].data_labels.number_format = "0%"
            prs.save(native)
            receipt = render(native, self.font(), "wechat-article", image)
            rp = root / "receipt.json"; rp.write_text(json.dumps(receipt))
            qa = {"editable_source": {"path": str(native), "sha256": receipt["source_sha256"],
                "rendered_image_sha256": receipt["final_png_sha256"], "render_receipt": str(rp)},
                "additional_text_review": [{"text": "Growth", "origin": "structural-label", "reason": "Series heading"}]}
            plan = {"rendering": {"strategy": "editable-chart-final"}, "visual_card": {"platform": "wechat-article"},
                    "structured_content": {"dataset": {"units": "%", "rows": [{"label": "Revenue", "value": 30}]}}}
            qp = root / "qa.json"; qp.write_text(json.dumps(qa))
            self.assertEqual(reviewed_editable_source(plan, qa, qp, receipt["final_png_sha256"], image)["re_render_pixel_difference_ae"], 0)
            plan["structured_content"]["dataset"]["rows"][0]["value"] = .3
            with self.assertRaisesRegex(ValueError, "dataset labels and values"):
                reviewed_editable_source(plan, qa, qp, receipt["final_png_sha256"], image)

if __name__ == "__main__":
    unittest.main()
